from pathlib import Path
import hashlib
import shutil
import sys

from openpyxl import load_workbook
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from rule_maintainer import RuleMaintainer
from workbuddy_bridge import WorkBuddyBridge


@pytest.fixture
def publication(tmp_path):
    (tmp_path / "config").mkdir()
    target = tmp_path / "config/rules_structured.xlsx"
    shutil.copyfile(ROOT / "config/rules_structured.xlsx", target)
    original = target.read_bytes()
    incoming = tmp_path / "reviewed.xlsx"
    book = load_workbook(target)
    book.create_sheet("publication_test")["A1"] = "reviewed test revision"
    book.save(incoming)
    book.close()
    return RuleMaintainer(tmp_path, {}), incoming, original, hashlib.sha256(original).hexdigest()


def test_reviewed_publication_preserves_backup_and_metadata(publication):
    maintainer, incoming, original, digest = publication
    outcome = maintainer.publish_reviewed_workbook(incoming, expected_sha256=digest, reviewer="tester", reason="confirmed test revision")
    assert maintainer.structured_rules_path.read_bytes() == incoming.read_bytes()
    assert Path(outcome["backup"]).read_bytes() == original
    assert Path(outcome["backup"]).with_suffix(".json").exists()


@pytest.mark.parametrize("invalid", ["format", "stale", "reviewer"])
def test_invalid_or_unreviewed_publication_keeps_original(publication, invalid):
    maintainer, incoming, original, digest = publication
    if invalid == "format":
        incoming.write_bytes(b"invalid workbook")
    with pytest.raises(ValueError):
        maintainer.publish_reviewed_workbook(incoming, expected_sha256="stale" if invalid == "stale" else digest, reviewer="" if invalid == "reviewer" else "tester", reason="test")
    assert maintainer.structured_rules_path.read_bytes() == original


def test_publication_is_blocked_during_workbuddy_processing(publication):
    maintainer, incoming, original, digest = publication
    with WorkBuddyBridge(maintainer.root_dir, {}).exclusive():
        with pytest.raises(RuntimeError, match="另一个"):
            maintainer.publish_reviewed_workbook(incoming, expected_sha256=digest, reviewer="tester", reason="test")
    assert maintainer.structured_rules_path.read_bytes() == original


def test_failed_replace_keeps_original(publication, monkeypatch):
    maintainer, incoming, original, digest = publication
    def fail(*args):
        raise PermissionError("file locked")
    monkeypatch.setattr("excel_exports.os.replace", fail)
    with pytest.raises(PermissionError):
        maintainer.publish_reviewed_workbook(incoming, expected_sha256=digest, reviewer="tester", reason="test")
    assert maintainer.structured_rules_path.read_bytes() == original
