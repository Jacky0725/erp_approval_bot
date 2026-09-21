"""Map review-only P0 rule candidates to historical regression samples."""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBENCH = max((ROOT / "outputs").glob("rule_governance_workbench_*"), key=lambda path: path.stat().st_mtime)
OUT_DIR = ROOT / "outputs" / f"rule_candidate_impact_{datetime.now():%Y%m%d_%H%M%S}"


def read_sheet(path: Path, sheet: str) -> list[dict[str, object]]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    data = ws.iter_rows(values_only=True)
    headers = [str(value or "") for value in next(data)]
    return [dict(zip(headers, row)) for row in data if any(value is not None for value in row)]


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook_path = next(WORKBENCH.glob("*.xlsx"))
    candidates = read_sheet(workbook_path, "规则修复候选")
    samples = read_sheet(workbook_path, "历史回归样本")
    impacts: list[dict[str, object]] = []
    summaries: list[dict[str, object]] = []
    for candidate in candidates:
        pattern = str(candidate.get("current_pattern") or "").strip()
        if not pattern:
            continue
        matcher = re.compile(re.escape(pattern), re.I)
        hits = 0
        for sample in samples:
            searchable = " ".join(str(sample.get(key) or "") for key in ("raw_name", "standard_name", "current_reason"))
            if not matcher.search(searchable):
                continue
            hits += 1
            impacts.append({
                "candidate_id": candidate["candidate_id"], "proposed_action": candidate["proposed_action"],
                "source_rule_id": candidate["source_rule_id"], "pattern": pattern,
                "list_number": sample["list_number"], "sequence": sample["sequence"],
                "raw_name": sample["raw_name"], "cas": sample["cas"],
                "current_category": sample["expected_category"], "current_reason": sample["current_reason"],
                "manual_review": sample["manual_review"],
                "proposed_outcome": "待业务确认；本报告不改变判定结果。",
            })
        summaries.append({"candidate_id": candidate["candidate_id"], "source_rule_id": candidate["source_rule_id"],
                          "pattern": pattern, "proposed_action": candidate["proposed_action"], "historical_hits": hits,
                          "risk": "高" if hits else "低", "approval_status": "待审核"})

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "影响概要"
    headers = list(summaries[0])
    ws.append(headers)
    for item in summaries: ws.append([item.get(key, "") for key in headers])
    detail_headers = list(impacts[0]) if impacts else ["candidate_id", "source_rule_id", "pattern"]
    detail = wb.create_sheet("样本影响明细")
    detail.append(detail_headers)
    for item in impacts: detail.append([item.get(key, "") for key in detail_headers])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(cell.value or "")) for cell in col) + 2))
    output = OUT_DIR / "P0规则候选影响报告.xlsx"
    wb.save(output)
    print(output)


if __name__ == "__main__": main()
