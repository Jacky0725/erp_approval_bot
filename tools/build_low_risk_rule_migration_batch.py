"""Prepare, but never apply, the first low-risk structured-rule migration batch."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
IMPACT_DIR = max((ROOT / "outputs").glob("rule_candidate_impact_*"), key=lambda path: path.stat().st_mtime)
WORKBENCH_DIR = max((ROOT / "outputs").glob("rule_governance_workbench_*"), key=lambda path: path.stat().st_mtime)
OUT_DIR = ROOT / "outputs" / f"rule_migration_batch_01_{datetime.now():%Y%m%d_%H%M%S}"

def load(path: Path, sheet: str) -> list[dict[str, object]]:
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True); ws=wb[sheet]
    it=ws.iter_rows(values_only=True); headers=[str(x or '') for x in next(it)]
    return [dict(zip(headers, row)) for row in it if any(x is not None for x in row)]

def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    impact=load(next(IMPACT_DIR.glob('*.xlsx')), '影响概要')
    candidates={row['candidate_id']: row for row in load(next(WORKBENCH_DIR.glob('*.xlsx')), '规则修复候选')}
    batch=[]
    for row in impact:
        candidate=candidates.get(row['candidate_id'], {})
        if row.get('historical_hits') != 0 or candidate.get('proposed_action') != '停用':
            continue
        batch.append({
            'candidate_id': row['candidate_id'], 'source_rule_id': row['source_rule_id'],
            'category': candidate.get('category'), 'pattern': row['pattern'],
            'proposed_change': '将 enabled 改为 false；不删除行，保留规则 ID 与审计历史。',
            'shadow_acceptance': '运行全部回归样本；该规则停用后不得产生类别差异。',
            'rollback': '将 enabled 恢复为 true。', 'approval_status': '待业务审核', 'approved_by': '',
        })
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='第一批停用候选'
    headers=list(batch[0]) if batch else ['candidate_id']; ws.append(headers)
    for row in batch: ws.append([row.get(h,'') for h in headers])
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(60,max(12,max(len(str(c.value or '')) for c in col)+2))
    out=OUT_DIR/'第一批低风险规则停用候选.xlsx'; wb.save(out)
    (OUT_DIR/'README.md').write_text('# 第一批低风险规则停用候选\n\n仅限影子回归和审核。未修改正式规则。批准后才允许按规则 ID 将 enabled 改为 false。\n',encoding='utf-8')
    print(out); print(f'candidates={len(batch)}')

if __name__ == '__main__': main()
