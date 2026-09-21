"""Read-only audit for the structured reagent rule workbook.

This tool never edits the production workbook or settings.  It produces an
auditable Excel report and JSON snapshot under outputs/ for review before a
rule migration is proposed.
"""
from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
RULES_PATH = ROOT / "config" / "rules_structured.xlsx"
OUT_DIR = ROOT / "outputs" / f"rule_governance_audit_{datetime.now():%Y%m%d_%H%M%S}"

FRAGMENT_TERMS = {"kg", "气体", "蒸汽", "毒性", "氧化性", "还原性", "5mg", "50mg"}


def rows(sheet: openpyxl.worksheet.worksheet.Worksheet) -> list[dict[str, object]]:
    values = list(sheet.iter_rows(values_only=True))
    headers = [str(value or "").strip() for value in values[0]]
    return [dict(zip(headers, value)) for value in values[1:] if any(item is not None for item in value)]


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(RULES_PATH, data_only=False)
    rule_rows = [row for row in rows(workbook["rules"]) if row.get("enabled") is not False]
    findings: list[dict[str, object]] = []

    by_pattern: dict[str, list[dict[str, object]]] = defaultdict(list)
    for rule in rule_rows:
        rule_id = str(rule.get("rule_id") or "").strip()
        pattern = str(rule.get("pattern") or "").strip()
        category = str(rule.get("category") or "").strip()
        condition = str(rule.get("condition") or "").strip()
        scope = str(rule.get("field_scope") or "").strip()
        by_pattern[pattern.casefold()].append(rule)
        if pattern.casefold() in FRAGMENT_TERMS or len(pattern) <= 2:
            findings.append({"severity": "P0", "action": "停用候选", "rule_id": rule_id,
                             "category": category, "pattern": pattern, "condition": condition,
                             "reason": "关键词过宽，属于说明文字碎片；应替换为结构化阈值或明确物质族规则。"})
        if rule_id.startswith(("LEGACY-ACU-", "LEGACY-TOX-", "LEGACY-SAC-")):
            findings.append({"severity": "P0", "action": "迁移候选", "rule_id": rule_id,
                             "category": category, "pattern": pattern, "condition": condition,
                             "reason": "历史说明文字拆分规则；应由 LD50/LC50、浓度或明确危险属性规则替代。"})
        if not scope or not condition:
            findings.append({"severity": "P1", "action": "补全字段", "rule_id": rule_id,
                             "category": category, "pattern": pattern, "condition": condition,
                             "reason": "缺少匹配范围或条件，审计和回归难以解释。"})

    for pattern, candidates in by_pattern.items():
        categories = {str(row.get("category") or "").strip() for row in candidates}
        if pattern and len(categories) > 1:
            findings.append({"severity": "P1", "action": "验证互斥", "rule_id": ", ".join(str(row.get("rule_id") or "") for row in candidates),
                             "category": ", ".join(sorted(categories)), "pattern": pattern,
                             "condition": " | ".join(str(row.get("condition") or "") for row in candidates),
                             "reason": "同一模式跨类别；必须证明条件互斥并添加边界回归样本。"})

    report = openpyxl.Workbook()
    overview = report.active
    overview.title = "概要"
    overview.append(["项目", "值"])
    overview.append(["规则文件", str(RULES_PATH)])
    overview.append(["规则指纹", "sha256:" + hashlib.sha256(RULES_PATH.read_bytes()).hexdigest()])
    overview.append(["生成时间", datetime.now().isoformat(timespec="seconds")])
    overview.append(["启用规则数", len(rule_rows)])
    overview.append(["发现项数", len(findings)])
    finding_sheet = report.create_sheet("治理发现")
    headers = ["severity", "action", "rule_id", "category", "pattern", "condition", "reason"]
    finding_sheet.append(headers)
    for finding in findings:
        finding_sheet.append([finding.get(header, "") for header in headers])
    inventory = report.create_sheet("规则清单")
    inventory_headers = ["rule_id", "category", "match_type", "field_scope", "pattern", "condition", "confidence", "description", "enabled"]
    inventory.append(inventory_headers)
    for rule in rule_rows:
        inventory.append([rule.get(header, "") for header in inventory_headers])
    for sheet in report.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    output = OUT_DIR / "规则治理审计.xlsx"
    report.save(output)
    (OUT_DIR / "governance_findings.json").write_text(json.dumps(findings, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT_DIR / "README.md").write_text(
        "# 规则治理审计\n\n本工件为只读审计结果；未修改 config/rules_structured.xlsx、V1 配置或 ERP 数据。\n"
        "下一步：逐项确认‘治理发现’，在规则候选表中提出替代规则，再运行历史回归后才允许启用。\n",
        encoding="utf-8",
    )
    print(output)


if __name__ == "__main__":
    main()
