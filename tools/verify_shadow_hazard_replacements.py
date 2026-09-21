"""Validate disabled shadow replacements with positive and negative SDS/GHS cases."""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
SHADOW=max((ROOT/'outputs').glob('rule_shadow_batch_01_*'),key=lambda p:p.stat().st_mtime)
OUT=ROOT/'outputs'/f'rule_shadow_hazard_validation_{datetime.now():%Y%m%d_%H%M%S}'
CASES=[
 ('冲击敏感正例','测试物','SDS: 对冲击敏感；对摩擦敏感。','易爆类'),
 ('摩擦敏感英文正例','test material','Safety data: friction-sensitive material.','易爆类'),
 ('H250正例','测试物','GHS Hazard statement H250: Catches fire spontaneously if exposed to air.','强反应性'),
 ('自热正例','测试物','SDS: self-heating substance.','强反应性'),
 ('普通碰撞对照','测试物','运输过程中避免碰撞和跌落。',''),
 ('普通自燃对照','测试物','请勿接近火源；本品不具有自燃危险。',''),
]
def main():
 OUT.mkdir(parents=True,exist_ok=True); source=next(SHADOW.glob('rules_structured_shadow_batch_01.xlsx')); target=OUT/'rules_structured_shadow_hazard_enabled.xlsx'; shutil.copy2(source,target)
 wb=openpyxl.load_workbook(target); ws=wb['rules']; headers=[c.value for c in ws[1]]; enabled=headers.index('enabled')+1
 for row in range(2,ws.max_row+1):
  if str(ws.cell(row=row,column=1).value or '').startswith('SHADOW-'): ws.cell(row=row,column=enabled).value=True
 # Tighten the shadow candidate before testing: a bare “自燃” also matches
 # negated SDS prose such as “不具有自燃危险”.
 for row in range(2, ws.max_row + 1):
  if ws.cell(row=row, column=1).value == 'SHADOW-REA-PYRO-001':
   ws.cell(row=row, column=5).value = r'(?:\bH\s*250\b|\bH\s*251\b|\bH\s*252\b|发生自燃|易自燃|自燃性|自燃物质|自热|pyrophoric|self[- ]heating)'
 wb.save(target)
 patterns={}
 for row in ws.iter_rows(min_row=2,values_only=True):
  if str(row[0] or '').startswith('SHADOW-'): patterns[str(row[0])]=str(row[4])
 results=[]
 for name,raw_name,text,expected in CASES:
  explosive=bool(__import__('re').search(patterns['SHADOW-EXP-IMPACT-001'],text,flags=__import__('re').I))
  reactive=bool(__import__('re').search(patterns['SHADOW-REA-PYRO-001'],text,flags=__import__('re').I))
  actual='易爆类' if explosive else ('强反应性' if reactive else '')
  ok=(actual==expected) if expected else not actual
  results.append((name,expected or '不应为易爆/强反应性',actual,ok,text))
 out=openpyxl.Workbook(); ws=out.active; ws.title='专项回归'; ws.append(['案例','预期','实际','通过','证据文本'])
 for result in results: ws.append(result)
 ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
 for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(80,max(12,max(len(str(c.value or '')) for c in col)+2))
 report=OUT/'危险属性替代规则专项回归.xlsx'; out.save(report); print(target); print(report); print('passed='+str(sum(row[3] for row in results))+'/'+str(len(results)))
if __name__=='__main__': main()
