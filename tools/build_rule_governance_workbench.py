"""Create review-only rule candidates and a historical regression workbook."""
from __future__ import annotations

import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
AUDIT_DIR = max((ROOT / "outputs").glob("rule_governance_audit_*"), key=lambda path: path.stat().st_mtime)
OUT_DIR = ROOT / "outputs" / f"rule_governance_workbench_{datetime.now():%Y%m%d_%H%M%S}"


def records(path: Path) -> list[dict[str, object]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    values = ws.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(values)]
    return [dict(zip(headers, row)) for row in values if any(value is not None for value in row)]


def write_sheet(workbook: openpyxl.Workbook, name: str, headers: list[str], data: list[dict[str, object]]) -> None:
    ws = workbook.create_sheet(name)
    ws.append(headers)
    for row in data:
        ws.append([row.get(header, "") for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    findings = json.loads((AUDIT_DIR / "governance_findings.json").read_text(encoding="utf-8"))
    candidates: list[dict[str, object]] = []
    for finding in findings:
        if finding["severity"] != "P0":
            continue
        action = "停用" if finding["action"] == "停用候选" else "替换"
        candidates.append({
            "candidate_id": f"RGC-{len(candidates)+1:03d}", "proposed_action": action,
            "source_rule_id": finding["rule_id"], "category": finding["category"],
            "current_pattern": finding["pattern"], "current_condition": finding["condition"],
            "proposed_rule": "以结构化属性/阈值或明确物质族规则替代" if action == "替换" else "停用该宽泛关键词规则",
            "required_regression": "命中、未命中、上下阈值、冲突类别各至少一例",
            "approval_status": "待审核", "reviewer": "", "review_note": finding["reason"],
        })

    suggestions = records(ROOT / "data/logs/approval_suggestions_all.xlsx")
    review_rows = records(ROOT / "data/review_queue.xlsx")
    regression: list[dict[str, object]] = []
    seen: set[tuple[str, str]] = set()
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in suggestions:
        category = str(row.get("最终建议类别") or "未分类")
        grouped[category].append(row)
    for category, rows in grouped.items():
        for row in rows[:25]:
            key = (str(row.get("试剂名称") or ""), str(row.get("CAS号") or ""))
            if not key[0] or key in seen:
                continue
            seen.add(key)
            regression.append({"source": "approval_suggestions", "list_number": row.get("试剂清单号"), "sequence": row.get("序号"),
                               "raw_name": key[0], "cas": key[1], "standard_name": row.get("标准化名称"),
                               "expected_category": category, "current_reason": row.get("规则原因"),
                               "manual_review": row.get("需人工复核"), "review_status": "待回归"})
    for row in review_rows:
        key = (str(row.get("试剂名称") or row.get("chemical_name") or ""), str(row.get("cas") or ""))
        if not key[0] or key in seen:
            continue
        seen.add(key)
        regression.append({"source": "manual_review", "list_number": row.get("试剂清单号"), "sequence": row.get("序号"),
                           "raw_name": key[0], "cas": key[1], "standard_name": row.get("standard_name"),
                           "expected_category": row.get("expected_category") or row.get("suggested_category"),
                           "current_reason": row.get("reason") or row.get("display_reason"),
                           "manual_review": True, "review_status": "待回归"})

    workbook = openpyxl.Workbook()
    overview = workbook.active
    overview.title = "说明"
    overview.append(["项目", "值"])
    overview.append(["生成时间", datetime.now().isoformat(timespec="seconds")])
    overview.append(["规则候选数", len(candidates)])
    overview.append(["回归样本数", len(regression)])
    overview.append(["安全状态", "仅供审核；不修改正式规则、不写 ERP"])
    write_sheet(workbook, "规则修复候选", list(candidates[0]), candidates)
    write_sheet(workbook, "历史回归样本", list(regression[0]), regression)
    for ws in workbook.worksheets:
        ws.freeze_panes = "A2"
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    output = OUT_DIR / "规则治理工作台.xlsx"
    workbook.save(output)
    (OUT_DIR / "README.md").write_text("# 规则治理工作台\n\n候选规则必须经过回归并由业务审核批准，才可进入正式规则表。\n", encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
