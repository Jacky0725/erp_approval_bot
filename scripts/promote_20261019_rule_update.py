"""Promote the selected A-F rows without restoring obsolete Unknown rules.

Usage: python scripts/promote_20261019_rule_update.py INCOMING_WORKBOOK
The existing project workbook is backed up before the atomic replacement.
"""

from __future__ import annotations

import argparse
from copy import copy
from datetime import datetime
from pathlib import Path
import shutil

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "config" / "rules_structured.xlsx"
RULE_PREFIXES = ("BUS-NOR-DRUG-", "BUS-NOR-BIO-", "BUS-NOR-POLYMER-")


def _header(sheet):
    return {str(cell.value or "").strip(): cell.column for cell in sheet[1]}


def _rows(sheet, key_columns):
    columns = _header(sheet)
    return {
        tuple(str(sheet.cell(row, columns[key]).value or "").strip() for key in key_columns): row
        for row in range(2, sheet.max_row + 1)
    }


def _append_row(source, source_row, target, overrides=None):
    overrides = overrides or {}
    source_columns = _header(source)
    target_columns = _header(target)
    row = target.max_row + 1
    for field, dest_column in target_columns.items():
        if field not in source_columns:
            continue
        source_cell = source.cell(source_row, source_columns[field])
        dest_cell = target.cell(row, dest_column)
        dest_cell.value = overrides.get(field, source_cell.value)
        if source_cell.has_style:
            dest_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            dest_cell.number_format = source_cell.number_format
    return row


def promote(incoming: Path) -> dict[str, int]:
    current = load_workbook(TARGET)
    source = load_workbook(incoming, read_only=False)
    changes = {"rules": 0, "examples": 0, "aliases": 0, "thresholds": 0, "notes": 0}

    source_rules = source["rules"]
    target_rules = current["rules"]
    source_columns = _header(source_rules)
    target_columns = _header(target_rules)
    target_rule_rows = _rows(target_rules, ("rule_id",))
    for row in range(2, source_rules.max_row + 1):
        rule_id = str(source_rules.cell(row, source_columns["rule_id"]).value or "").strip()
        if not rule_id.startswith(RULE_PREFIXES) or (rule_id,) in target_rule_rows:
            continue
        # The source workbook expresses family precedence in prose. The
        # engine already evaluates all matches by category priority, so an
        # executable `any` condition is used. Drug identity is not permitted
        # to mask a separately established higher-priority hazard.
        _append_row(source_rules, row, target_rules, {"condition": "any"})
        changes["rules"] += 1
    policy_rule_id = "BUS-NOR-CLEAR-NAME-NO-EVIDENCE-001"
    policy_values = {
        "category": "普通类", "match_type": "business_policy", "field_scope": "name",
        "pattern": "explicit chemical ERP name; no usable web/local hazard evidence",
        "condition": "V1 clear-name no-evidence policy; executed in approval_flow",
        "confidence": 0.8, "enabled": True,
        "description": "仅 V1：明确化学原名在无可用来源证据且本地危险规则无结论时，按用户确认的业务口径归普通类；不是无危险性的证明。",
    }
    if (policy_rule_id,) not in target_rule_rows:
        policy_row = target_rules.max_row + 1
        target_rules.cell(policy_row, target_columns["rule_id"]).value = policy_rule_id
        for field, value in policy_values.items():
            if field in target_columns:
                target_rules.cell(policy_row, target_columns[field]).value = value
        changes["rules"] += 1

    # F: the source's free-text condition is not executable. Use the existing
    # liquid-context gate, which also accepts H224/H225/H226 themselves as
    # explicit flammable-liquid evidence after source relevance checking.
    flammable_row = target_rule_rows.get(("BUS-FLA-GHS-H224-226-001",))
    if flammable_row:
        target_rules.cell(flammable_row, target_columns["condition"]).value = "requires_liquid_context"
        changes["rules"] += 1

    for sheet_name, keys in (
        ("examples", ("category", "example_name", "match_mode")),
        ("aliases", ("alias", "standard_name", "cas")),
    ):
        src, dst = source[sheet_name], current[sheet_name]
        src_columns = _header(src)
        seen = set(_rows(dst, keys))
        for row in range(2, src.max_row + 1):
            key = tuple(str(src.cell(row, src_columns[field]).value or "").strip() for field in keys)
            if not any(key) or key in seen:
                continue
            if sheet_name == "examples" and key[0] != "普通类":
                continue
            _append_row(src, row, dst)
            seen.add(key)
            changes[sheet_name] += 1

    src_thresholds, dst_thresholds = source["thresholds"], current["thresholds"]
    threshold_row = _rows(dst_thresholds, ("threshold_id",)).get(("T-TOX-002",))
    incoming_row = _rows(src_thresholds, ("threshold_id",)).get(("T-TOX-002",))
    if threshold_row and incoming_row:
        for field in ("value", "description"):
            dst_thresholds.cell(threshold_row, _header(dst_thresholds)[field]).value = src_thresholds.cell(
                incoming_row, _header(src_thresholds)[field]
            ).value
        changes["thresholds"] += 1

    if "notes" in source.sheetnames and "notes" in current.sheetnames:
        src, dst = source["notes"], current["notes"]
        seen = set(_rows(dst, ("note_id",)))
        columns = _header(src)
        for row in range(2, src.max_row + 1):
            note_id = str(src.cell(row, columns["note_id"]).value or "").strip()
            if not note_id or (note_id,) in seen or not note_id.endswith("20261015"):
                continue
            # The uploaded 'final override' note is not an executable rule and
            # would falsely imply that a hazardous drug becomes Ordinary.
            if note_id == "NOTE-DRUG-OVERRIDE-20261015":
                continue
            _append_row(src, row, dst)
            changes["notes"] += 1

    backup_dir = ROOT / "config" / "rule_backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"rules_structured.before_selected_20261019_{datetime.now():%H%M%S}.xlsx"
    shutil.copy2(TARGET, backup)
    staged = TARGET.with_suffix(".xlsx.staged")
    current.save(staged)
    # Validate the staged file before replacing the live rules.
    import sys
    sys.path.insert(0, str(ROOT / "src"))
    from rule_workbook_validator import validate_rule_workbook
    validation = validate_rule_workbook(staged)
    if not validation.valid:
        staged.unlink(missing_ok=True)
        raise RuntimeError(f"Promoted workbook failed validation: {validation.errors}")
    staged.replace(TARGET)
    changes["backup"] = str(backup)
    return changes


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("incoming", type=Path)
    args = parser.parse_args()
    print(promote(args.incoming))
