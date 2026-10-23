"""Small, resumable file hand-off to the WorkBuddy reagent project.

No model calls or chemical decisions live here. ERP identity comes exclusively
from the saved input, never from an ID invented in the returned spreadsheet.
"""
from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
from typing import Any, Iterable, Iterator
from uuid import uuid4

import pandas as pd
from openpyxl import load_workbook

from audit_logger import AuditLogger
from category_mapper import erp_candidates_for_rule_category, erp_property_options
from excel_exports import write_excel_atomic


RESULT_COLUMNS = ["批次ID", "记录ID", "试剂名称", "物化特性类别", "判断依据"]
REVIEW_CATEGORIES = {"", "需人工复核", "无法判定", "证据不足"}
EMPTY_PROPERTIES = {"", "-", "—"}


def text(value: Any) -> str:
    return "" if value is None or value is False else str(value).strip()


def record_identity(record: dict[str, Any]) -> dict[str, str]:
    return {
        "试剂清单号": text(record.get("试剂清单号") or record.get("清单号")),
        "_erp_record_id": text(record.get("_erp_record_id") or record.get("id")),
        "试剂名称": text(record.get("试剂名称") or record.get("name")),
        "CAS号": text(record.get("CAS号") or record.get("cas_code")),
        "序号": text(record.get("序号") or record.get("sequence")),
        "type": text(record.get("type")),
        "desc": text(record.get("desc")),
    }


def record_key(record: dict[str, Any]) -> str:
    return hashlib.sha256(json.dumps(record_identity(record), ensure_ascii=False, sort_keys=True).encode()).hexdigest()


class WorkBuddyBridge:
    def __init__(self, root_dir: Path, settings: dict[str, Any]):
        self.root_dir = Path(root_dir)
        self.settings = settings
        self.config = settings.get("workbuddy", {}) or {}
        self.directory = self.root_dir / self.config.get("batch_dir", "data/workbuddy")
        self.audit = AuditLogger.from_settings(settings, self.root_dir)

    def rules_path(self) -> Path:
        path = Path(self.settings.get("paths", {}).get("structured_rules_excel", "config/rules_structured.xlsx"))
        return path if path.is_absolute() else self.root_dir / path

    @contextmanager
    def exclusive(self) -> Iterator[None]:
        """OS lock is released even when a worker is killed; no stale lock flag."""
        self.directory.mkdir(parents=True, exist_ok=True)
        with (self.directory / ".lock").open("a+b") as stream:
            try:
                # Windows byte-range locks also prohibit reads from another
                # handle. Inspect length, not the byte another worker locked.
                if stream.seek(0, os.SEEK_END) == 0:
                    stream.write(b"0")
                    stream.flush()
                stream.seek(0)
                if os.name == "nt":
                    import msvcrt
                    msvcrt.locking(stream.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl
                    fcntl.flock(stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
            except OSError as error:
                raise RuntimeError("另一个 WorkBuddy 批次正在处理，请勿重复启动。") from error
            try:
                yield
            finally:
                stream.seek(0)
                if os.name == "nt":
                    import msvcrt
                    msvcrt.locking(stream.fileno(), msvcrt.LK_UNLCK, 1)
                else:
                    import fcntl
                    fcntl.flock(stream, fcntl.LOCK_UN)

    def batches(self) -> list[dict[str, Any]]:
        batches = []
        for path in sorted(self.directory.glob("*/state.json")):
            state = json.loads(path.read_text(encoding="utf-8"))
            if path.parent.name != state["batch_id"]:
                raise ValueError(f"批次目录与状态不一致：{path}")
            batches.append(state)
        return batches

    def batch_path(self, state: dict[str, Any]) -> Path:
        batch_id = state["batch_id"]
        if not batch_id or Path(batch_id).name != batch_id or any(c not in "0123456789_abcdef" for c in batch_id):
            raise ValueError("批次 ID 无效")
        path = (self.directory / batch_id).resolve()
        if path.parent != self.directory.resolve():
            raise ValueError("批次路径越界")
        return path

    def save(self, state: dict[str, Any]) -> None:
        path = self.batch_path(state) / "state.json"
        state["updated_at"] = datetime.now().isoformat(timespec="seconds")
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)

    def prepare(self, records: Iterable[dict[str, Any]]) -> dict[str, Any] | None:
        mode = self.config.get("prompt_mode", "full")
        if mode not in {"full", "project"}:
            raise ValueError("workbuddy.prompt_mode 必须为 full 或 project。")
        rule_bytes = self.rules_path().read_bytes()
        rule_digest = hashlib.sha256(rule_bytes).hexdigest()
        existing = self.batches()
        waiting = [s for s in existing if s["status"] == "waiting_result" and s["rules_sha256"] == rule_digest]
        if waiting:
            return waiting[0]
        # A failed row is retried from its original batch, not silently reissued.
        known = {item["key"] for s in existing if s["rules_sha256"] == rule_digest for item in s["items"]}
        stable_ids: set[tuple[str, str]] = set()
        selected = []
        size = max(1, min(20, int(self.config.get("batch_size", 20))))
        for record in records:
            identity = record_identity(record)
            if not identity["_erp_record_id"] or not identity["试剂清单号"] or not identity["试剂名称"]:
                raise ValueError("WorkBuddy 输入缺少稳定明细 ID、清单号或试剂名称。")
            if record_key(record) in known or text(record.get("物化特性")) not in EMPTY_PROPERTIES:
                continue
            stable_id = (identity["试剂清单号"], identity["_erp_record_id"])
            if stable_id in stable_ids:
                raise ValueError("输入中同一 ERP 明细 ID 对应多条不同记录。")
            stable_ids.add(stable_id)
            known.add(record_key(record))
            selected.append({"key": record_key(record), "input": identity, "status": "pending", "attempts": 0})
            if len(selected) == size:
                break
        if not selected:
            return None
        batch_id = datetime.now().strftime("%Y%m%d_%H%M%S_") + uuid4().hex[:8]
        state = {"batch_id": batch_id, "status": "waiting_result", "rules_sha256": hashlib.sha256(rule_bytes).hexdigest(), "items": selected}
        path = self.batch_path(state)
        path.mkdir(parents=True)
        (path / "rules.xlsx").write_bytes(rule_bytes)
        rows = [{**item["input"], "批次ID": batch_id, "记录ID": item["key"]} for item in selected]
        write_excel_atomic(pd.DataFrame(rows), path / "input.xlsx")
        result_rows = [{"批次ID": batch_id, "记录ID": i["key"], "试剂名称": i["input"]["试剂名称"], "物化特性类别": "", "判断依据": ""} for i in selected]
        write_excel_atomic(pd.DataFrame(result_rows, columns=RESULT_COLUMNS), path / "result_template.xlsx")
        prompt = (
            f"试剂审批批次 {batch_id}，共 {len(selected)} 条。独立判定，不访问或写入 ERP，不修改规则。\n"
            f"输入：{path / 'input.xlsx'}\n规则：{path / 'rules.xlsx'}\n"
            f"请完整读取本批规则（含优先级及排除条件），以此快照为准。\n"
            "优先复用已有 Excel 处理工具或可用 Python 环境，避免每批新建环境或重复安装依赖。\n"
            f"填写模板：{path / 'result_template.xlsx'}，仅一个工作表 Sheet1，严格五列：{'、'.join(RESULT_COLUMNS)}。\n"
            "批次ID、记录ID、试剂名称原样返回，每条一行，不漏项、不重复，不使用公式。"
            "无法判定时类别填“需人工复核”，判断依据说明缺失信息；不得猜测。"
            "判断依据写明适用规则和事实来源，模型知识如实标注。\n"
            f"先完整保存并关闭 {path / 'result.pending.xlsx'}，最后重命名为 {path / 'result.xlsx'}。"
            "只有正式文件保存成功后才回复完成。不要只返回表格或代码。"
            "请在 pending 文件上完成校验，重命名必须是最后一次文件操作；发布后只回复完成，不再读取、修改或重建本批文件。"
        )
        if mode == "project":
            prompt = f"按 WB-EXCEL-v1 执行批次目录：{path}。仅处理本批，结果保存到该目录的 result.xlsx。"
        state["prompt_mode"] = mode
        (path / "prompt.txt").write_text(prompt, encoding="utf-8")
        self.save(state)
        self.audit.info(json.dumps({"event": "workbuddy_batch_prepared", "batch_id": batch_id, "rows": len(selected)}, ensure_ascii=False))
        return state

    def load_results(self, state: dict[str, Any]) -> None:
        rules = self.batch_path(state) / "rules.xlsx"
        if hashlib.sha256(rules.read_bytes()).hexdigest() != state["rules_sha256"]:
            raise ValueError("本批规则快照被修改，禁止导入。")
        path = self.batch_path(state) / "result.xlsx"
        if path.resolve().parent != self.batch_path(state):
            raise ValueError("结果文件路径越界。")
        payload = path.read_bytes()
        digest = hashlib.sha256(payload).hexdigest()
        if state.get("result_sha256"):
            if digest != state["result_sha256"]:
                raise ValueError("已导入的结果文件发生变化；为避免同批混用，拒绝继续写入。")
            return
        # Read formulas as formulas: data_only would conceal them as cached values.
        import io
        workbook = load_workbook(io.BytesIO(payload), read_only=False, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                raise ValueError("结果 Excel 必须只有一个工作表。")
            sheet = workbook.worksheets[0]
            if sheet.merged_cells.ranges:
                raise ValueError("结果 Excel 不允许合并单元格。")
            if sheet.max_row > 21 or sheet.max_column != len(RESULT_COLUMNS):
                raise ValueError("结果 Excel 行列数不符合本批约定。")
            if any(c.data_type == "f" for row in sheet for c in row):
                raise ValueError("结果 Excel 不允许公式。")
            rows = list(sheet.iter_rows(values_only=True))
            if not rows or list(rows[0]) != RESULT_COLUMNS:
                raise ValueError(f"结果表头必须为：{'、'.join(RESULT_COLUMNS)}")
            expected = {i["key"]: i for i in state["items"]}
            parsed = {}
            for raw in rows[1:]:
                if not any(v is not None for v in raw):
                    continue
                values = dict(zip(RESULT_COLUMNS, map(text, raw)))
                key = values["记录ID"]
                if key not in expected or key in parsed or values["批次ID"] != state["batch_id"]:
                    raise ValueError("结果存在重复记录、未知记录或批次不一致。")
                if values["试剂名称"] != expected[key]["input"]["试剂名称"] or not values["判断依据"]:
                    raise ValueError("试剂名称被修改或缺少判断依据。")
                parsed[key] = values
            if set(parsed) != set(expected):
                raise ValueError("结果记录不完整，禁止部分导入。")
            for key, values in parsed.items():
                expected[key]["result"] = values
            state["result_sha256"] = digest
            state["status"] = "imported"
            self.save(state)
        finally:
            workbook.close()

    def process(self, state: dict[str, Any], client: Any, memory: Any, *, dry_run: bool = True) -> dict[str, Any]:
        """Reconcile the original record before each write; no chemical reclassification."""
        if state["status"] == "completed":
            return state
        self.load_results(state)
        if not dry_run and hashlib.sha256(self.rules_path().read_bytes()).hexdigest() != state["rules_sha256"]:
            raise ValueError("正式规则已更新，本批旧版本结果禁止写入；请按新规则重新判定。")
        for item in state["items"]:
            if item["status"] in {"saved", "review"}:
                continue
            result = item["result"]
            category = result["物化特性类别"]
            candidates = [c for c in erp_candidates_for_rule_category(category, self.settings) if c in erp_property_options(self.settings)]
            if dry_run:
                print(f"[workbuddy preview] {item['input']['试剂名称']} -> {category}; ERP={candidates}")
                continue
            # Store model provenance separately from manually verified memory.
            memory.record_workbuddy_result(state["batch_id"], item["key"], item["input"], result)
            if category in REVIEW_CATEGORIES:
                self._record(state, item, "review", "判定", "MODEL_UNRESOLVED", result["判断依据"])
                continue
            if len(candidates) != 1:
                self._record(state, item, "blocked", "类别映射", "CATEGORY_MAPPING", f"判定类别“{category}”没有唯一 ERP 映射；候选：{candidates}")
                continue
            target = candidates[0]
            item["erp_category"] = target
            identity = item["input"]
            try:
                records = client.fetch_reagent_detail(identity["试剂清单号"])
                matches = [r for r in records if text(r.get("_erp_record_id") or r.get("id")) == identity["_erp_record_id"]]
                if len(matches) != 1 or record_identity(matches[0]) != identity:
                    self._record(state, item, "blocked", "写前核对", "IDENTITY_CHANGED", "ERP 明细身份或判定相关信息已变化。")
                    continue
                if matches[0].get("_erp_parent_state") != "wait":
                    self._record(state, item, "blocked", "写前核对", "TASK_NOT_PENDING", "ERP 清单不是待审批状态，或未能读取其状态。")
                    continue
                if "物化特性" not in matches[0]:
                    self._record(state, item, "blocked", "写前核对", "PROPERTY_UNAVAILABLE", "ERP 返回中缺少物化特性字段，未写入。")
                    continue
                current = text(matches[0].get("物化特性"))
                if current == target:
                    if item.get("write_started"):
                        memory.record_workbuddy_result(state["batch_id"], item["key"], identity, result, synced=True)
                        self._record(state, item, "saved", "回读", "", "上次写入结果已在 ERP 中确认。")
                    else:
                        self._record(state, item, "blocked", "写前核对", "ALREADY_CLASSIFIED", "ERP 已有类别，未覆盖；不能确认由本批写入。")
                    continue
                if current not in EMPTY_PROPERTIES:
                    self._record(state, item, "blocked", "写前核对", "PROPERTY_CHANGED", f"ERP 物化特性已为“{current}”，未覆盖。")
                    continue
                if item["attempts"] >= max(1, int(self.config.get("max_write_attempts", 2))):
                    continue
                item["attempts"] += 1
                item["write_started"] = True
                self.save(state)  # Durable before sending a non-idempotent request.
                outcome = client.save_physicochemical_property_by_id(identity["_erp_record_id"], target, {**identity, "最终建议类别": target}, matches[0])
                if outcome.saved and outcome.verified:
                    memory.record_workbuddy_result(state["batch_id"], item["key"], identity, result, synced=True)
                    self._record(state, item, "saved", "回读", "", outcome.detail)
                else:
                    code = "ERP_CATEGORY_UNAVAILABLE" if "category is not available" in outcome.detail else "ERP_WRITE"
                    self._record(state, item, "failed", "保存/回读", code, outcome.detail)
            except Exception as error:
                self._record(state, item, "failed", "保存/回读", "ERP_ERROR", str(error))
        if not dry_run:
            state["status"] = "completed" if all(i["status"] == "saved" for i in state["items"]) else "needs_attention"
            self.save(state)
            if state["status"] == "completed":
                self.cleanup(state)
            else:
                count = sum(i["status"] != "saved" for i in state["items"])
                print(f"[workbuddy attention] 批次 {state['batch_id']} 有 {count} 条未完成，结果文件已保留。")
        return state

    def cleanup(self, state: dict[str, Any]) -> None:
        if state["status"] != "completed" or not all(i["status"] == "saved" for i in state["items"]):
            return
        try:
            (self.batch_path(state) / "result.xlsx").unlink(missing_ok=True)
            state["cleaned"] = True
            state.pop("cleanup_error", None)
        except OSError as error:
            state["cleanup_error"] = str(error)
        self.save(state)

    def _record(self, state: dict[str, Any], item: dict[str, Any], status: str, stage: str, code: str, reason: str) -> None:
        item.update(status=status, stage=stage, error_code=code, detail=reason)
        event = {"event": "workbuddy_result", "batch_id": state["batch_id"], "record_id": item["key"], "status": status, "stage": stage, "code": code, "detail": reason, "category": item["result"]["物化特性类别"], "erp_category": item.get("erp_category", "")}
        self.audit.info(json.dumps(event, ensure_ascii=False))
        self.save(state)
        print(f"[workbuddy] {item['input']['试剂名称']} {status}: {reason}")
